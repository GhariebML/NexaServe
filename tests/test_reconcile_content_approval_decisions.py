import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from scripts.reconcile_content_approval_decisions import (
    APPROVED,
    REQUIRED_HEADERS,
    REVIEW_SHEETS,
    read_workbook_decisions,
    reconcile_decisions,
)


DATASET = "FAQs/converted/depi_helper.json"
RECORD_ID = "synthetic_record_1"
CONTENT_HASH = "a" * 64
FILE_HASH = "b" * 64
QUESTION_AR = "كيف أستعيد كلمة المرور؟"
ANSWER_EN = "Use the official password reset page."
LOCATOR = {"source_sheet": "R5", "source_row": 7}


def fixture_pair():
    queue_record = {
        "source_json": DATASET,
        "source_record_id": RECORD_ID,
        "original_content_hash": CONTENT_HASH,
        "source_original_file": "FAQs/DEPI HELPER.xlsx",
        "source_locator": LOCATOR,
        "original_question": QUESTION_AR,
        "original_answer": ANSWER_EN,
        "language": "Arabic/English",
        "recommended_disposition": "NEEDS_OWNER_APPROVAL",
    }
    current_record = {
        "content_hash": CONTENT_HASH,
        "original_source_file": queue_record["source_original_file"],
        "source_file_hash_matches": True,
        "question": QUESTION_AR,
        "answer": ANSWER_EN,
        "language": "Arabic/English",
        "source_locator": LOCATOR,
    }
    row = {
        "Stable Source Record ID": RECORD_ID,
        "Source JSON Path": DATASET,
        "Original Source File": queue_record["source_original_file"],
        "Original Source Locator": json.dumps(LOCATOR, ensure_ascii=False, sort_keys=True),
        "Source URL": None,
        "Content Hash": CONTENT_HASH,
        "Question": QUESTION_AR,
        "Answer": ANSWER_EN,
        "Language": "Arabic/English",
        "Current Program Label": "DEPI",
        "Reviewed Program Identity": "DEPI",
        "Proposed Program": "DEPI",
        "Proposed Retrieval Scope": "DEPI_SCOPED",
        "Source Authority Classification": "internal DEPI guidance",
        "Current Review Status": "NEEDS_OWNER_APPROVAL",
        "Exact Review Rationale": "Synthetic test fixture",
        "Missing Provenance / Scope Evidence": "Owner approval",
        "Evidence Required to Resolve": "Owner approval reference",
        "Duplicate / Relationship IDs": None,
        "Reviewer Decision": "APPROVED_INTERNAL_DEPI_GUIDANCE",
        "Reviewer Identity / Responsible Role": "Reviewer identity: Synthetic Reviewer; Role: Authorized DEPI Content Owner",
        "Decision Date (UTC)": datetime(2026, 10, 12, 11, 23, 14),
        "Approval Evidence / Decision Reference": "OWNER_APPROVAL:SYNTHETIC-001",
        "Reviewer Rationale": "Synthetic explicit approval for unit testing.",
        "_sheet": "Owner Approval Queue",
        "_row": 5,
    }
    return row, queue_record, {(DATASET, RECORD_ID): [current_record]}


def reconcile(rows, queue, current):
    return reconcile_decisions(rows, queue, current, APPROVED | {
        "OUT_OF_SCOPE", "REJECTED", "NEEDS_MORE_EVIDENCE", "NEEDS_OWNER_APPROVAL"
    })


class ReconcileContentApprovalDecisionsTests(unittest.TestCase):
    def test_valid_authorized_internal_approval(self):
        row, source, current = fixture_pair()
        result = reconcile([row], [source], current)
        self.assertEqual(result["approved_record_count"], 1)
        self.assertEqual(result["ledger_entries"][0]["decision"], "APPROVED_INTERNAL_DEPI_GUIDANCE")
        self.assertEqual(result["ledger_entries"][0]["reviewer_identity"], "Synthetic Reviewer")

    def test_missing_reviewer_identity_or_evidence_is_withheld(self):
        row, source, current = fixture_pair()
        row["Reviewer Identity / Responsible Role"] = "Role: Authorized DEPI Content Owner"
        row["Approval Evidence / Decision Reference"] = None
        result = reconcile([row], [source], current)
        self.assertEqual(result["approved_record_count"], 0)
        self.assertEqual(len(result["incomplete_decisions"]), 1)
        self.assertTrue(any("Reviewer identity" in x for x in result["incomplete_decisions"][0]["missing_fields"]))
        self.assertIn("Approval Evidence / Decision Reference", result["incomplete_decisions"][0]["missing_fields"])

    def test_invalid_decision_status_is_rejected(self):
        row, source, current = fixture_pair()
        row["Reviewer Decision"] = "ACCEPTED"
        result = reconcile([row], [source], current)
        self.assertEqual(result["ledger_entries"], [])
        self.assertEqual(len(result["invalid_decisions_or_source_cells"]), 1)

    def test_duplicate_record_decisions_are_withheld(self):
        row, source, current = fixture_pair()
        duplicate = dict(row)
        duplicate["_row"] = 6
        duplicate["Reviewer Decision"] = "REJECTED"
        result = reconcile([row, duplicate], [source], current)
        self.assertEqual(result["ledger_entries"], [])
        self.assertEqual(len(result["conflicting_or_duplicated_decisions"]), 1)

    def test_unknown_source_record_is_not_written(self):
        row, source, current = fixture_pair()
        row["Stable Source Record ID"] = "unknown_synthetic_id"
        result = reconcile([row], [source], current)
        self.assertEqual(result["ledger_entries"], [])
        self.assertEqual(len(result["unknown_source_records"]), 1)

    def test_changed_current_hash_marks_approval_stale(self):
        row, source, current = fixture_pair()
        current[(DATASET, RECORD_ID)][0]["content_hash"] = "c" * 64
        result = reconcile([row], [source], current)
        self.assertEqual(result["ledger_entries"], [])
        self.assertEqual(len(result["stale_source_hashes"]), 1)

    def test_empty_decisions_remain_unreviewed(self):
        row, source, current = fixture_pair()
        row["Reviewer Decision"] = None
        for field in (
            "Reviewer Identity / Responsible Role", "Decision Date (UTC)",
            "Approval Evidence / Decision Reference", "Reviewer Rationale",
        ):
            row[field] = None
        result = reconcile([row], [source], current)
        self.assertEqual(result["approved_record_count"], 0)
        self.assertEqual(result["unreviewed_rows"], 1)
        self.assertEqual(result["ledger_entries"], [])

    def test_reviewer_name_without_explicit_decision_does_not_approve(self):
        row, source, current = fixture_pair()
        row["Reviewer Decision"] = None
        row["Reviewer Identity / Responsible Role"] = "Reviewer identity: Synthetic Reviewer; Role: Authorized DEPI Content Owner"
        row["Decision Date (UTC)"] = None
        row["Approval Evidence / Decision Reference"] = None
        row["Reviewer Rationale"] = None
        result = reconcile([row], [source], current)
        self.assertEqual(result["approved_record_count"], 0)
        self.assertEqual(result["ledger_entries"], [])
        self.assertEqual(len(result["incomplete_decisions"]), 1)

    def test_arabic_and_english_cells_preserved_by_workbook_reader(self):
        row, _, _ = fixture_pair()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "synthetic_review.xlsx"
            workbook = Workbook()
            workbook.remove(workbook.active)
            for sheet_name in REVIEW_SHEETS:
                worksheet = workbook.create_sheet(sheet_name)
                for column, header in enumerate(REQUIRED_HEADERS, start=1):
                    worksheet.cell(4, column, header)
                if sheet_name == "Owner Approval Queue":
                    for column, header in enumerate(REQUIRED_HEADERS, start=1):
                        worksheet.cell(5, column, row.get(header))
            workbook.save(path)
            parsed, errors = read_workbook_decisions(path)
        self.assertEqual(errors, [])
        self.assertEqual(parsed[0]["Question"], QUESTION_AR)
        self.assertEqual(parsed[0]["Answer"], ANSWER_EN)

    def test_reconciliation_is_idempotent(self):
        row, source, current = fixture_pair()
        first = reconcile([row], [source], current)
        second = reconcile([row], [source], current)
        self.assertEqual(first, second)
        self.assertEqual(
            json.dumps(first["ledger_entries"], ensure_ascii=False, sort_keys=True),
            json.dumps(second["ledger_entries"], ensure_ascii=False, sort_keys=True),
        )


if __name__ == "__main__":
    unittest.main()
