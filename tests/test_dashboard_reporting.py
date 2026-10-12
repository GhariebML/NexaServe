"""Isolated tests for dashboard report generation and evaluation summaries."""
import asyncio
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook
from fastapi import FastAPI
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[1]
os.environ.setdefault("CS_DB_PASSWORD", "test-only-not-a-real-secret")
os.environ.setdefault("DASHBOARD_JWT_SECRET", "test-only-dashboard-signing-key")
sys.path.insert(0, str(ROOT / "dashboard" / "backend"))

import routes
import database
from evaluation_summary import (
    build_sanitized_evaluation_artifact,
    summarize_latest_rag_evaluation,
    summarize_stored_rag_evaluation,
)
from reporting import build_workbook, safe_cell_value, safe_filename


def _summary_artifacts(tmp_path, *, passed=2, failed=1, blocked=1, runtime=None, corpus=None):
    master = tmp_path / "master.json"
    result = tmp_path / "rag.json"
    master.write_text(json.dumps({
        "finished": "2026-10-11T14:00:00",
        "verdict": "NOT_READY",
        "suite_results": [{"test": "Golden RAG Suite", "status": "PASS" if failed == 0 and blocked == 0 else "FAIL", "stdout_tail": "Full report saved to: tests/golden_test_results.json"}],
    }), encoding="utf-8")
    summary = {
        "total": passed + failed + blocked,
        "passed": passed,
        "failed": failed,
        "blocked": blocked,
        "timestamp": "2026-10-11T13:59:00",
    }
    artifact = {"summary": summary, "results": [
        {"status": status, "answer": "must not be returned"}
        for status, count in (("PASS", passed), ("FAIL", failed), ("BLOCKED", blocked))
        for _ in range(count)
    ]}
    if runtime:
        artifact["runtime_identity"] = runtime
    if corpus:
        artifact["corpus_version"] = corpus
    result.write_text(json.dumps(artifact, ensure_ascii=False), encoding="utf-8")
    return master, result


def test_workbook_stylish_rtl_and_preserves_arabic_english(tmp_path):
    content = build_workbook(
        "ملخص NexaServe",
        [{
            "name": "المحتوى",
            "title": "English and العربية",
            "columns": [
                {"key": "english", "label": "English", "type": "text"},
                {"key": "arabic", "label": "العربية", "type": "text"},
                {"key": "when", "label": "Timestamp", "type": "datetime"},
            ],
            "rows": [{"english": "DEPI FAQ", "arabic": "مبادرة رواد مصر الرقمية", "when": "2026-10-12T10:30:00+02:00"}],
        }],
        filters={"program": "DEPI"}, lang="ar",
        generated_at=datetime(2026, 10, 12, 9, 0, tzinfo=timezone.utc),
    )
    path = tmp_path / "report.xlsx"
    path.write_bytes(content)
    workbook = load_workbook(path, data_only=False)
    ws = workbook["المحتوى"]
    assert ws.sheet_view.rightToLeft is True
    assert ws.freeze_panes == "A6"
    assert ws.auto_filter.ref == "A5:C6"
    assert ws["A6"].value == "DEPI FAQ"
    assert ws["B6"].value == "مبادرة رواد مصر الرقمية"
    assert ws["C6"].value == datetime(2026, 10, 12, 8, 30)
    assert ws["A5"].fill.fgColor.rgb.endswith("2563EB")


@pytest.mark.parametrize("input_value", ["=1+1", " +SUM(A1:A2)", "-cmd|' /C calc'!A0", "@SUM(A1)"])
def test_formula_injection_is_written_as_inert_text(input_value):
    assert safe_cell_value(input_value).startswith("'")
    assert safe_cell_value(-4) == -4


def test_filename_is_safe_and_uses_supplied_current_date():
    assert safe_filename("Customers & Students", datetime(2026, 10, 12).date()) == "NexaServe_Customers_Students_Report_2026-10-12.xlsx"


def test_evaluation_summary_is_blocked_and_does_not_expose_benchmark_answers(tmp_path):
    master, result = _summary_artifacts(tmp_path)
    summary = summarize_latest_rag_evaluation(master, result)
    assert summary["status"] == "BLOCKED"
    assert summary["passed"] == 2 and summary["failed"] == 1 and summary["blocked"] == 1
    assert summary["evaluated_cases"] == 3 and summary["planned_cases"] == 4
    assert summary["pass_rate"] == 0.6667
    assert summary["historical"] is True
    assert "answer" not in json.dumps(summary)


def test_evaluation_is_only_linked_when_runtime_and_corpus_match(tmp_path):
    master, result = _summary_artifacts(tmp_path, runtime="candidate-x", corpus="kb-7")
    assert summarize_latest_rag_evaluation(master, result, "other", "kb-7")["status"] == "BLOCKED"
    assert summarize_latest_rag_evaluation(master, result, "candidate-x", "kb-7")["status"] == "NEEDS_IMPROVEMENT"
    clean_master = tmp_path / "clean-master.json"
    clean_artifact = tmp_path / "clean-rag.json"
    clean_master.write_text(json.dumps({"verdict": "READY", "finished": "2026-10-12T10:00:00+00:00", "suite_results": [{"test": "Golden RAG Suite", "status": "PASS", "stdout_tail": "Full report saved to: tests/golden_test_results.json"}]}), encoding="utf-8")
    clean_artifact.write_text(json.dumps({"runtime_identity": "candidate-x", "corpus_version": "kb-7", "summary": {"total": 2, "passed": 2, "failed": 0, "blocked": 0, "timestamp": "2026-10-12T10:00:00+00:00"}}), encoding="utf-8")
    result = summarize_latest_rag_evaluation(clean_master, clean_artifact, "candidate-x", "kb-7")
    assert result["status"] == "HEALTHY"
    assert result["historical"] is False


def test_runtime_linked_old_evaluation_is_stale(tmp_path):
    master, result = _summary_artifacts(tmp_path, passed=2, failed=0, blocked=0, runtime="candidate-x", corpus="kb-7")
    data = json.loads(result.read_text(encoding="utf-8"))
    data["summary"]["timestamp"] = "2026-10-01T10:00:00+00:00"
    result.write_text(json.dumps(data), encoding="utf-8")
    master_data = json.loads(master.read_text(encoding="utf-8"))
    master_data["finished"] = "2026-10-01T10:00:00+00:00"
    master.write_text(json.dumps(master_data), encoding="utf-8")
    summary = summarize_latest_rag_evaluation(
        master, result, "candidate-x", "kb-7", now=datetime(2026, 10, 12, tzinfo=timezone.utc)
    )
    assert summary["status"] == "STALE"


def test_missing_evaluation_is_not_evaluated(tmp_path):
    result = summarize_latest_rag_evaluation(tmp_path / "absent-master.json", tmp_path / "absent-rag.json")
    assert result["status"] == "NOT_EVALUATED"
    assert result["pass_rate"] is None


def test_packaged_evaluation_artifact_contains_only_sanitized_summary(tmp_path):
    master, result = _summary_artifacts(tmp_path)
    artifact = build_sanitized_evaluation_artifact(master, result)
    packaged = tmp_path / "rag_evaluation_summary.json"
    packaged.write_text(json.dumps(artifact, ensure_ascii=False), encoding="utf-8")
    summary = summarize_stored_rag_evaluation(packaged)
    assert summary["status"] == "BLOCKED"
    assert summary["passed"] == 2 and summary["failed"] == 1 and summary["blocked"] == 1
    assert "answer" not in packaged.read_text(encoding="utf-8")
    assert len(summary["source_artifact_sha256"]) == 64


def test_customer_report_exports_all_filtered_pages_with_role_masking(monkeypatch):
    async def fake_get_customers(*, search, program, page, limit, user):
        assert search == "اسم" and program == "DEPI" and user["role"] == "agent"
        rows = [{"id": f"c{index}",
                 "full_name": "=HYPERLINK(\"x\")" if index == 1 else ("ليلى" if index == 0 else f"Customer {index}"),
                 "phone_masked": "010****123", "program": "DEPI",
                 "created_at": "2026-10-01T10:00:00+00:00", "last_interaction": None,
                 "conversation_count": 2, "open_tickets": 0, "customer_status": "Active"}
                for index in range(101)]
        start = (page - 1) * limit
        return {"items": rows[start:start + limit], "total": 101, "page": page, "limit": limit, "pages": 2}

    monkeypatch.setattr(routes, "get_customers", fake_get_customers)
    response = asyncio.run(routes.export_dashboard_report(
        report_type="customers", period=None, program="DEPI", search="اسم",
        ticket_status=None, priority=None, sla_breached=None,
        execution_status=None, lang="ar", user={"role": "agent"},
    ))
    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "NexaServe_customers_Report_" in response.headers["content-disposition"]
    async def read_stream():
        return b"".join([chunk async for chunk in response.body_iterator])
    workbook = load_workbook(__import__("io").BytesIO(asyncio.run(read_stream())))
    assert workbook["Customers"].max_row == 106  # all 101 filtered rows, not only the first UI page
    assert workbook["Customers"]["B7"].value.startswith("'")


def test_empty_export_decision_rows_create_a_valid_empty_sheet():
    data = build_workbook("Empty", [{"name": "Rows", "columns": [{"key": "id", "label": "ID", "type": "text"}], "rows": []}])
    workbook = load_workbook(__import__("io").BytesIO(data))
    assert workbook["Rows"].max_row == 5
    assert workbook["Rows"].auto_filter.ref == "A5:A5"


def test_export_query_timeout_is_scoped_and_parameterized(monkeypatch):
    executed = []

    class Cursor:
        description = [("value",)]
        def __enter__(self): return self
        def __exit__(self, *_args): return False
        def execute(self, sql, params=()): executed.append((sql, params))
        def fetchall(self): return [(1,)]

    class Connection:
        def cursor(self): return Cursor()

    from contextlib import contextmanager
    @contextmanager
    def fake_connection():
        yield Connection()

    monkeypatch.setattr(database, "get_db", fake_connection)
    token = database.REPORT_QUERY_DEADLINE.set(__import__("time").monotonic() + 5)
    try:
        assert database.query_all("SELECT 1 AS value") == [{"value": 1}]
    finally:
        database.REPORT_QUERY_DEADLINE.reset(token)
    assert executed[0][0] == "SELECT set_config('statement_timeout', %s, true)"
    assert executed[0][1][0].endswith("ms")
    assert executed[1] == ("SELECT 1 AS value", ())


def test_dashboard_frontend_exposes_bilingual_shared_export_and_historical_label():
    app = (ROOT / "dashboard/frontend/app.js").read_text(encoding="utf-8")
    html = (ROOT / "dashboard/frontend/index.html").read_text(encoding="utf-8")
    assert "export_report: \"Export to Excel\"" in app
    assert "export_report: \"تصدير إلى Excel\"" in app
    assert "exportCurrentReport()" in html
    assert "Acceptance test pass rate among completed cases. This is not answer accuracy." in app
    assert "loadRAGEvaluationCard()" in app
    export_map = app.split("const EXPORT_REPORTS = {", 1)[1].split("};", 1)[0]
    assert "health: 'health'" in export_map and "simulator:" not in export_map


def test_evaluation_and_export_routes_require_auth_and_work_without_database(monkeypatch):
    app = FastAPI()
    app.include_router(routes.router)
    client = TestClient(app)
    assert client.get("/api/dashboard/rag/evaluation-summary").status_code == 401
    assert client.get("/api/dashboard/reports/executive/excel").status_code == 401
    app.dependency_overrides[routes.get_current_user] = lambda: {"role": "readonly"}
    evaluation = client.get("/api/dashboard/rag/evaluation-summary")
    assert evaluation.status_code == 200
    evaluation_json = evaluation.json()
    assert evaluation_json["status"] == "BLOCKED"
    assert (evaluation_json["passed"], evaluation_json["failed"], evaluation_json["blocked"]) == (33, 2, 1)
    assert "answer" not in evaluation.text

    async def fake_summary(*, period, program, user):
        assert period == "7d" and program == "DEPI" and user["role"] == "readonly"
        return {"total_customers": 4, "rag_success_rate": None}

    monkeypatch.setattr(routes, "get_summary", fake_summary)
    response = client.get("/api/dashboard/reports/executive/excel?period=7d&program=DEPI")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    workbook = load_workbook(__import__("io").BytesIO(response.content))
    assert workbook["Overview"]["B6"].value == 4
    assert "period: 7d" in workbook["Overview"]["B3"].value


def test_each_passive_dashboard_page_has_a_page_specific_workbook(monkeypatch):
    async def summary(**_kwargs):
        assert _kwargs["period"] == "7d" and _kwargs["program"] == "DEPI"
        return {"total_customers": 4}

    async def customers(**_kwargs):
        assert _kwargs["program"] == "DEPI" and _kwargs["search"] == ""
        return {"items": [{"id": "c-1", "full_name": "ليلى", "program": "DEPI", "phone_masked": "010****123", "created_at": None, "last_interaction": None, "conversation_count": 1, "open_tickets": 0, "customer_status": "Active"}], "total": 1, "pages": 1}

    async def conversations(**_kwargs):
        assert _kwargs["period"] == "7d" and _kwargs["program"] == "DEPI"
        return {"items": [{"id": "v-1", "customer_name": "ليلى", "program": "DEPI", "channel": "web", "intent": "faq", "avg_confidence": .8, "status": "closed", "created_at": None, "message_count": 2}], "total": 1, "pages": 1, "metrics": {"total_conversations": 1}, "charts": {}}

    async def tickets(**_kwargs):
        assert _kwargs["status"] == "open" and _kwargs["priority"] == "high" and _kwargs["program"] == "DEPI"
        return {"items": [{"ticket_number": "T-1", "customer_name": "ليلى", "program": "DEPI", "priority": "high", "status": "open", "assigned_agent": "agent", "sla_deadline": None, "sla_status": "ON_TRACK", "created_at": None, "updated_at": None, "reason": "private reason"}], "total": 1, "pages": 1}

    async def rag(**_kwargs):
        return {"rag_queries": 5, "confidence_distribution": [{"bracket": "High", "count": 3}], "program_isolation": {"status": "not_measured"}}

    async def evaluation(**_kwargs):
        return {"status": "BLOCKED", "pass_rate": None}

    async def knowledge(**_kwargs):
        return {"total_kb_documents": 1, "documents_by_program": [{"program": "DEPI", "count": 1}], "documents_by_source": [{"source": "official", "count": 1}]}

    async def health(**_kwargs):
        return {"services": [{"name": "Database", "status": "ONLINE", "latency_ms": 4, "last_check": "2026-10-12T10:00:00+00:00", "error": "not exported"}]}

    async def activity(**_kwargs):
        return {"items": [{"id": "a-1", "event_type": "turn_completed", "workflow_name": "wf", "channel": "web", "latency_ms": 10, "created_at": None, "description": "not exported"}]}

    async def executions(**_kwargs):
        assert _kwargs["status_filter"] == "failed"
        return {"items": []}

    async def workflows(**_kwargs):
        return {"items": [{"id": "wf-1", "name": "Workflow", "nodes_count": 2, "tags": ["safe"], "active": False, "updated_at": None}]}

    async def n8n_executions(**_kwargs):
        return {"items": [{"id": "exec-1", "workflow_id": "wf-1", "status": "success", "mode": "manual", "finished": True, "started_at": None, "stopped_at": None, "duration_ms": 5}]}

    for name, fn in {
        "get_summary": summary, "get_customers": customers, "get_conversations": conversations,
        "get_tickets": tickets, "get_rag_metrics": rag, "get_rag_evaluation_summary": evaluation,
        "get_kb_metrics": knowledge, "get_system_health": health,
        "get_recent_activity": activity, "list_executions": executions,
        "get_n8n_workflows": workflows, "get_n8n_executions": n8n_executions,
    }.items():
        monkeypatch.setattr(routes, name, fn)

    expected_sheets = {
        "executive": {"Overview"}, "customers": {"Customers"},
        "conversations": {"Conversations", "Analytics"}, "tickets": {"Tickets"},
        "rag": {"Runtime telemetry", "Confidence", "Evaluation"},
        "knowledge-base": {"Summary", "By program", "By source"},
        "health": {"Health"}, "activity": {"Recent requests", "Audit activity"},
        "n8n": {"Workflows", "Latest executions"},
    }
    ticket_workbook = None
    for report_type, names in expected_sheets.items():
        response = asyncio.run(routes.export_dashboard_report(
            report_type=report_type, period="7d", program="DEPI", search="",
            ticket_status="open", priority="high", sla_breached=None,
            execution_status="failed", lang="en", user={"role": "agent"},
        ))
        async def read_response():
            return b"".join([chunk async for chunk in response.body_iterator])
        workbook = load_workbook(__import__("io").BytesIO(asyncio.run(read_response())))
        assert names.issubset(set(workbook.sheetnames))
        if report_type == "tickets":
            ticket_workbook = workbook
    assert ticket_workbook["Tickets"].max_row == 6
    assert all(cell.value != "private reason" for row in ticket_workbook["Tickets"].iter_rows() for cell in row)
