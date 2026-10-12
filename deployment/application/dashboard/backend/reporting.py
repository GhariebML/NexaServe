"""Small, read-only Excel reporting helpers for the dashboard API."""
from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO
import re
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_EXPORT_ROWS = 5000


def safe_cell_value(value: Any) -> Any:
    """Keep untrusted text inert in spreadsheet applications."""
    if isinstance(value, str):
        stripped = value.lstrip(" \t\r\n\x00\x0b")
        if stripped.startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def safe_filename(report_name: str, generated_on: date | None = None) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", report_name).strip("_") or "Report"
    return f"NexaServe_{slug}_Report_{(generated_on or date.today()).isoformat()}.xlsx"


def _parse_datetime(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if parsed.tzinfo is not None:
            return parsed.astimezone(timezone.utc).replace(tzinfo=None)
        return parsed
    except ValueError:
        return value


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", " ", name).strip() or "Report"
    cleaned = cleaned[:31]
    base, candidate, index = cleaned, cleaned, 2
    while candidate.casefold() in used:
        suffix = f" {index}"
        candidate = f"{base[:31-len(suffix)]}{suffix}"
        index += 1
    used.add(candidate.casefold())
    return candidate


def _write_table(
    ws: Worksheet,
    title: str,
    columns: Sequence[Mapping[str, str]],
    rows: Iterable[Mapping[str, Any]],
    filters: Mapping[str, Any],
    generated_at: datetime,
    rtl: bool,
) -> None:
    ws.sheet_view.rightToLeft = rtl
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"
    navy = "142B4A"
    blue = "2563EB"
    muted = "E8EEF6"
    ws.cell(1, 1, safe_cell_value(title))
    ws.cell(1, 1).font = Font(name="Aptos", size=16, bold=True, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=navy)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.cell(2, 1, safe_cell_value("Generated at (UTC)"))
    ws.cell(2, 2, generated_at.astimezone(timezone.utc).replace(tzinfo=None))
    ws.cell(2, 2).number_format = "yyyy-mm-dd hh:mm:ss"
    ws.cell(3, 1, safe_cell_value("Active filters"))
    ws.cell(3, 2, safe_cell_value("; ".join(f"{k}: {v}" for k, v in filters.items()) or "None"))
    ws.cell(4, 1, safe_cell_value("Rows"))
    ws.cell(4, 2, 0)

    header_row = 5
    for index, column in enumerate(columns, 1):
        cell = ws.cell(header_row, index, safe_cell_value(column["label"]))
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 28

    row_count = 0
    for row_count, record in enumerate(rows, 1):
        if row_count > MAX_EXPORT_ROWS:
            raise ValueError(f"Report exceeds the {MAX_EXPORT_ROWS} row export limit")
        excel_row = header_row + row_count
        fill = "FFFFFF" if row_count % 2 else "F2F6FB"
        for col_index, column in enumerate(columns, 1):
            value = record.get(column["key"])
            if column.get("type") == "datetime":
                value = _parse_datetime(value)
            cell = ws.cell(excel_row, col_index, safe_cell_value(value))
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column.get("type") in {"text", "datetime"},
            )
            kind = column.get("type")
            if kind == "datetime" and isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif kind == "percent":
                cell.number_format = "0.0%"
            elif kind == "integer":
                cell.number_format = "#,##0"
            elif kind == "decimal":
                cell.number_format = "#,##0.00"
        ws.row_dimensions[excel_row].height = 30

    ws.cell(4, 2, row_count)
    if columns:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{header_row + row_count}"
    for col_index, column in enumerate(columns, 1):
        width = min(max(len(str(column["label"])) + 3, 14), 42)
        if column.get("type") == "text":
            width = min(max(width, 24), 48)
        ws.column_dimensions[get_column_letter(col_index)].width = width
    ws.column_dimensions["A"].width = min(max(ws.column_dimensions["A"].width or 14, 16), 40)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:5"


def build_workbook(
    report_title: str,
    sheets: Sequence[Mapping[str, Any]],
    filters: Mapping[str, Any] | None = None,
    lang: str = "en",
    generated_at: datetime | None = None,
) -> bytes:
    """Build a styled XLSX from already-authorized, already-filtered report rows."""
    if not sheets:
        raise ValueError("At least one report sheet is required")
    generated_at = generated_at or datetime.now().astimezone()
    filters = filters or {}
    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for index, spec in enumerate(sheets):
        name = _safe_sheet_name(str(spec.get("name") or f"Report {index+1}"), used)
        ws = workbook.create_sheet(name)
        _write_table(
            ws,
            str(spec.get("title") or report_title),
            spec["columns"],
            spec.get("rows", []),
            filters,
            generated_at,
            rtl=(lang == "ar"),
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
