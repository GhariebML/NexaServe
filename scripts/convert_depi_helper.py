#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
╔══════════════════════════════════════════════════════════════════════╗
║  NexaServe – DEPI HELPER.xlsx → JSON Professional Converter        ║
║  Converts multi-sheet Excel FAQ data into a clean, validated,      ║
║  RAG-ready JSON dataset with full provenance tracking.             ║
╚══════════════════════════════════════════════════════════════════════╝

Outputs:
  FAQs/converted/depi_helper.json      – Pretty-printed canonical JSON
  FAQs/converted/depi_helper.jsonl     – Line-delimited JSON (streaming)
  FAQs/converted/depi_helper_manifest.json – Conversion manifest & stats
"""

import hashlib
import json
import os
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

# ─── Configuration ────────────────────────────────────────────────────────────

SOURCE_FILE = Path(r"E:\NexaServe\FAQs\DEPI HELPER.xlsx")
OUTPUT_DIR  = Path(r"E:\NexaServe\FAQs\converted")

PROGRAM_SCOPE = "DEPI"
PROGRAM_FULL  = "Digital Egypt Pioneers Initiative"

# Schema for each sheet type
SHEET_PROFILES = {
    # ── Q&A sheets (col1=category/label, col2=response text) ──
    "ROUND 5": {
        "type": "qa_template",
        "round": "R5",
        "col_question": 1,
        "col_answer": 2,
        "category_prefix": "Round 5",
        "description": "Round 5 registration & onboarding FAQ responses",
    },
    "ROUND 4": {
        "type": "qa_template",
        "round": "R4",
        "col_question": 1,
        "col_answer": 2,
        "category_prefix": "Round 4",
        "description": "Round 4 registration & status FAQ responses",
    },
    "Round 5 Cancel": {
        "type": "qa_template",
        "round": "R5",
        "col_question": 1,
        "col_answer": 2,
        "category_prefix": "Round 5 – Cancellation",
        "description": "Round 5 cancellation-related responses",
    },
    "R4 CANCEL ": {
        "type": "qa_template",
        "round": "R4",
        "col_question": 1,
        "col_answer": 2,
        "category_prefix": "Round 4 – Cancellation",
        "description": "Round 4 cancellation & exam scheduling responses",
    },
    # ── Chatbot Q&A (ASK/ANSWER format) ──
    "chatting": {
        "type": "chatbot_qa",
        "round": "general",
        "col_question": 1,
        "col_answer": 2,
        "category_prefix": "Chatbot",
        "description": "Chatbot-style Q&A pairs for quick replies",
    },
    # ── R3 (header-only, may have hidden data) ──
    "R3": {
        "type": "qa_template",
        "round": "R3",
        "col_question": 1,
        "col_answer": 2,
        "category_prefix": "Round 3",
        "description": "Round 3 coaching & program FAQ responses",
    },
    # ── Email classification sheet (multi-column, #/التصنيف/نص الارسال) ──
    "تصنيف الميل المرحلة 3  ": {
        "type": "email_classification",
        "round": "R3",
        "col_index": 1,
        "col_category": 2,
        "col_message": 3,
        "col_approval": 4,
        "category_prefix": "R3 – Email Classification",
        "description": "Round 3 email response classification & templates",
    },
    # ── Email templates (complaint type / subject / message) ──
    "email ": {
        "type": "email_template",
        "round": "R3",
        "col_complaint_type": 1,
        "col_subject": 2,
        "col_message": 3,
        "category_prefix": "Email Templates",
        "description": "Email response templates by complaint type",
    },
}


# ─── Utility Functions ────────────────────────────────────────────────────────

def sha256_file(path: Path) -> str:
    """Compute SHA-256 hash of a file."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def sha256_text(text: str) -> str:
    """Compute SHA-256 hash of a text string."""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def clean_text(value: Any) -> str | None:
    """Normalize and clean a cell value."""
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() in ("none", "n/a", "nan", "-", ""):
        return None
    # Normalize Unicode
    text = unicodedata.normalize("NFC", text)
    # Collapse multiple whitespace but preserve newlines
    text = re.sub(r"[^\S\n]+", " ", text)
    # Collapse 3+ newlines into 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def detect_language(text: str | None) -> str:
    """Detect if text is primarily Arabic, English, or bilingual."""
    if not text:
        return "unknown"
    arabic_chars = len(re.findall(r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]", text))
    latin_chars = len(re.findall(r"[a-zA-Z]", text))
    total = arabic_chars + latin_chars
    if total == 0:
        return "unknown"
    ar_ratio = arabic_chars / total
    if ar_ratio > 0.7:
        return "ar"
    elif ar_ratio < 0.3:
        return "en"
    else:
        return "bilingual"


def infer_category(question: str | None, answer: str | None, sheet_profile: dict) -> str:
    """Infer a semantic category from the content."""
    text = (question or "") + " " + (answer or "")
    text_lower = text.lower()

    category_map = {
        "registration": ["تسجيل", "registration", "register", "تقديم", "التسجيل"],
        "cancellation": ["اعتذار", "cancel", "إلغاء", "الغاء", "انسحاب", "withdrawn"],
        "exam": ["امتحان", "exam", "اختبار", "batch slot", "test"],
        "track_change": ["تعديل التراك", "track", "تغيير", "تعديل الملف"],
        "documents": ["مستندات", "document", "إثبات", "شهادة", "أوراق"],
        "attendance": ["حضور", "attendance", "غياب", "absence"],
        "coaching": ["coaching", "كوتشينج", "جلسات"],
        "lms": ["lms", "منصة", "platform", "حساب", "account", "login", "دخول"],
        "results": ["نتيجة", "result", "نتائج", "قبول", "رفض", "accept", "reject"],
        "complaint": ["شكوى", "complaint", "مشكلة", "problem"],
        "data_edit": ["تعديل بيانات", "تعديل البيانات", "data edit", "update"],
        "transportation": ["تنقل", "مواصلات", "transport", "العاصمه الاداريه"],
        "certificates": ["شهادة", "certificate", "شهادات"],
        "schedule": ["موعد", "جدول", "schedule", "timeline"],
    }

    for cat, keywords in category_map.items():
        if any(kw in text_lower for kw in keywords):
            return cat

    return "general"


# ─── Sheet Extractors ─────────────────────────────────────────────────────────

def extract_qa_template(ws, profile: dict, sheet_name: str) -> list[dict]:
    """Extract Q&A template sheets (ROUND 5, ROUND 4, R4 CANCEL, etc.)."""
    records = []
    col_q = profile["col_question"]
    col_a = profile["col_answer"]

    for row in range(2, ws.max_row + 1):
        q_raw = clean_text(ws.cell(row, col_q).value)
        a_raw = clean_text(ws.cell(row, col_a).value)

        if not q_raw and not a_raw:
            continue

        question = q_raw or ""
        answer = a_raw or ""
        content = f"{question}\n---\n{answer}" if question and answer else (question or answer)

        records.append({
            "question": question if question else None,
            "answer": answer if answer else None,
            "content": content,
            "source_row": row,
        })

    return records


def extract_chatbot_qa(ws, profile: dict, sheet_name: str) -> list[dict]:
    """Extract chatbot-style ASK/ANSWER pairs."""
    records = []
    col_q = profile["col_question"]
    col_a = profile["col_answer"]

    for row in range(2, ws.max_row + 1):
        q_raw = clean_text(ws.cell(row, col_q).value)
        a_raw = clean_text(ws.cell(row, col_a).value)

        if not q_raw and not a_raw:
            continue

        question = q_raw or ""
        answer = a_raw or ""
        content = f"Q: {question}\nA: {answer}" if question and answer else (question or answer)

        records.append({
            "question": question if question else None,
            "answer": answer if answer else None,
            "content": content,
            "source_row": row,
        })

    return records


def extract_email_classification(ws, profile: dict, sheet_name: str) -> list[dict]:
    """Extract email classification sheet with index, category, message, approval."""
    records = []
    col_idx = profile["col_index"]
    col_cat = profile["col_category"]
    col_msg = profile["col_message"]
    col_appr = profile["col_approval"]

    for row in range(2, ws.max_row + 1):
        idx_raw = ws.cell(row, col_idx).value
        cat_raw = clean_text(ws.cell(row, col_cat).value)
        msg_raw = clean_text(ws.cell(row, col_msg).value)
        appr_raw = clean_text(ws.cell(row, col_appr).value)

        if not cat_raw and not msg_raw:
            continue

        category = cat_raw or "uncategorized"
        message = msg_raw or ""
        approval_status = appr_raw.upper() if appr_raw else None

        content_parts = [f"Category: {category}"]
        if message:
            content_parts.append(f"Message: {message}")
        if approval_status:
            content_parts.append(f"Status: {approval_status}")
        content = "\n".join(content_parts)

        records.append({
            "question": category,
            "answer": message,
            "content": content,
            "email_index": int(idx_raw) if idx_raw else None,
            "approval_status": approval_status,
            "source_row": row,
        })

    return records


def extract_email_template(ws, profile: dict, sheet_name: str) -> list[dict]:
    """Extract email template sheet (complaint type / subject / message)."""
    records = []
    col_type = profile["col_complaint_type"]
    col_subj = profile["col_subject"]
    col_msg = profile["col_message"]

    last_complaint_type = None
    for row in range(2, ws.max_row + 1):
        type_raw = clean_text(ws.cell(row, col_type).value)
        subj_raw = clean_text(ws.cell(row, col_subj).value)
        msg_raw = clean_text(ws.cell(row, col_msg).value)

        if not type_raw and not msg_raw:
            continue

        # Carry forward complaint type for grouped rows
        if type_raw:
            last_complaint_type = type_raw
        complaint_type = last_complaint_type or "general"
        subject = subj_raw or ""
        message = msg_raw or ""

        content_parts = [f"Complaint Type: {complaint_type}"]
        if subject:
            content_parts.append(f"Subject: {subject}")
        if message:
            content_parts.append(f"Message:\n{message}")
        content = "\n".join(content_parts)

        records.append({
            "question": complaint_type,
            "answer": message,
            "content": content,
            "email_subject": subject if subject else None,
            "complaint_type": complaint_type,
            "source_row": row,
        })

    return records


# ─── Extractor Dispatch ───────────────────────────────────────────────────────

EXTRACTORS = {
    "qa_template": extract_qa_template,
    "chatbot_qa": extract_chatbot_qa,
    "email_classification": extract_email_classification,
    "email_template": extract_email_template,
}


# ─── Main Pipeline ────────────────────────────────────────────────────────────

def main():
    print("╔══════════════════════════════════════════════════════════════╗")
    print("║  DEPI HELPER.xlsx → JSON Professional Converter            ║")
    print("╚══════════════════════════════════════════════════════════════╝")
    print()

    if not SOURCE_FILE.exists():
        print(f"❌ Source file not found: {SOURCE_FILE}")
        sys.exit(1)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # ── Step 1: Compute source file hash ──────────────────────────────────
    print("📋 Step 1: Computing source file hash...")
    source_hash = sha256_file(SOURCE_FILE)
    source_size = SOURCE_FILE.stat().st_size
    print(f"   SHA-256: {source_hash[:16]}...")
    print(f"   Size:    {source_size:,} bytes")
    print()

    # ── Step 2: Load workbook ─────────────────────────────────────────────
    print("📖 Step 2: Loading workbook...")
    wb = openpyxl.load_workbook(SOURCE_FILE, data_only=True)
    print(f"   Sheets:  {len(wb.sheetnames)}")
    for s in wb.sheetnames:
        print(f"     • {s}")
    print()

    # ── Step 3: Extract records from each sheet ───────────────────────────
    print("🔄 Step 3: Extracting records...")
    all_records = []
    sheet_stats = {}
    record_counter = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        profile = SHEET_PROFILES.get(sheet_name)

        if not profile:
            print(f"   ⚠️  Sheet '{sheet_name}' has no profile – skipping")
            sheet_stats[sheet_name] = {"status": "skipped", "records": 0}
            continue

        extractor = EXTRACTORS.get(profile["type"])
        if not extractor:
            print(f"   ⚠️  No extractor for type '{profile['type']}' – skipping")
            sheet_stats[sheet_name] = {"status": "no_extractor", "records": 0}
            continue

        raw_records = extractor(ws, profile, sheet_name)

        # Enrich each record with metadata
        for rec in raw_records:
            record_counter += 1
            content = rec.get("content", "")
            lang = detect_language(content)
            category = infer_category(rec.get("question"), rec.get("answer"), profile)

            enriched = {
                "id": f"depi_{record_counter:04d}",
                "program_scope": PROGRAM_SCOPE,
                "program": PROGRAM_FULL,
                "round": profile["round"],
                "sheet_type": profile["type"],
                "category": category,
                "category_label": profile["category_prefix"],
                "question": rec.get("question"),
                "answer": rec.get("answer"),
                "content": content,
                "language": lang,
                "source_file": str(SOURCE_FILE.name),
                "source_sheet": sheet_name.strip(),
                "source_row": rec.get("source_row"),
                "content_sha256": sha256_text(content),
            }

            # Add optional fields
            for extra_key in ("email_index", "approval_status", "email_subject", "complaint_type"):
                if extra_key in rec:
                    enriched[extra_key] = rec[extra_key]

            # RAG compatibility aliases
            enriched["question_ar"] = enriched["question"] if lang in ("ar", "bilingual") else None
            enriched["answer_ar"] = enriched["answer"] if lang in ("ar", "bilingual") else None
            enriched["source_attribution"] = f"{SOURCE_FILE.name} → {sheet_name.strip()} → Row {rec.get('source_row', '?')}"
            enriched["content_hash"] = enriched["content_sha256"]

            all_records.append(enriched)

        count = len(raw_records)
        sheet_stats[sheet_name.strip()] = {
            "status": "extracted",
            "records": count,
            "type": profile["type"],
            "description": profile["description"],
        }
        print(f"   ✅ {sheet_name.strip():30s} → {count:3d} records")

    print(f"\n   📊 Total raw records: {len(all_records)}")
    print()

    # ── Step 4: Deduplication ─────────────────────────────────────────────
    print("🔍 Step 4: Detecting duplicates...")
    seen_hashes = {}
    duplicates = []
    unique_records = []

    for rec in all_records:
        h = rec["content_sha256"]
        if h in seen_hashes:
            rec["is_duplicate"] = True
            rec["duplicate_of"] = seen_hashes[h]["id"]
            duplicates.append({
                "record_id": rec["id"],
                "duplicate_of": seen_hashes[h]["id"],
                "content_preview": rec["content"][:80],
                "source": rec["source_attribution"],
            })
        else:
            rec["is_duplicate"] = False
            rec["duplicate_of"] = None
            seen_hashes[h] = rec

        unique_records.append(rec)  # Keep all records, flag duplicates

    unique_count = sum(1 for r in unique_records if not r["is_duplicate"])
    dup_count = len(duplicates)
    print(f"   Unique:     {unique_count}")
    print(f"   Duplicates: {dup_count} (flagged, not removed)")
    print()

    # ── Step 5: Language distribution ─────────────────────────────────────
    print("🌐 Step 5: Language distribution...")
    lang_dist = {}
    for rec in unique_records:
        lang = rec["language"]
        lang_dist[lang] = lang_dist.get(lang, 0) + 1
    for lang, count in sorted(lang_dist.items(), key=lambda x: -x[1]):
        pct = (count / len(unique_records)) * 100
        print(f"   {lang:12s}: {count:4d} ({pct:.1f}%)")
    print()

    # ── Step 6: Category distribution ─────────────────────────────────────
    print("📂 Step 6: Category distribution...")
    cat_dist = {}
    for rec in unique_records:
        cat = rec["category"]
        cat_dist[cat] = cat_dist.get(cat, 0) + 1
    for cat, count in sorted(cat_dist.items(), key=lambda x: -x[1]):
        print(f"   {cat:25s}: {count:4d}")
    print()

    # ── Step 7: Write output files ────────────────────────────────────────
    generated_at = datetime.now(timezone.utc).isoformat()

    # 7a. Main JSON
    json_path = OUTPUT_DIR / "depi_helper.json"
    print(f"💾 Step 7a: Writing {json_path.name}...")
    output_json = {
        "schema_version": "2.0",
        "dataset_type": "faq_knowledge_base",
        "program_scope": PROGRAM_SCOPE,
        "program_name": PROGRAM_FULL,
        "generated_at": generated_at,
        "generator": "convert_depi_helper.py",
        "source_file": str(SOURCE_FILE.name),
        "source_file_sha256": source_hash,
        "source_file_size_bytes": source_size,
        "total_records": len(unique_records),
        "unique_records": unique_count,
        "duplicate_records": dup_count,
        "language_distribution": lang_dist,
        "category_distribution": cat_dist,
        "sheet_statistics": sheet_stats,
        "records": unique_records,
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(output_json, f, ensure_ascii=False, indent=2)
    print(f"   ✅ {json_path.stat().st_size:,} bytes")

    # 7b. JSONL (streaming format)
    jsonl_path = OUTPUT_DIR / "depi_helper.jsonl"
    print(f"💾 Step 7b: Writing {jsonl_path.name}...")
    with open(jsonl_path, "w", encoding="utf-8") as f:
        for rec in unique_records:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    print(f"   ✅ {jsonl_path.stat().st_size:,} bytes")

    # 7c. Manifest
    manifest_path = OUTPUT_DIR / "depi_helper_manifest.json"
    print(f"💾 Step 7c: Writing {manifest_path.name}...")
    manifest = {
        "pipeline_version": "2.0",
        "generated_at": generated_at,
        "generator_script": "scripts/convert_depi_helper.py",
        "source_file": str(SOURCE_FILE.name),
        "source_sha256": source_hash,
        "size_bytes": source_size,
        "sheets_processed": len([s for s in sheet_stats.values() if s.get("status") == "extracted"]),
        "sheets_skipped": len([s for s in sheet_stats.values() if s.get("status") != "extracted"]),
        "total_records": len(unique_records),
        "unique_records": unique_count,
        "duplicate_records": dup_count,
        "language_distribution": lang_dist,
        "category_distribution": cat_dist,
        "sheet_statistics": sheet_stats,
        "duplicates_detail": duplicates,
        "artifacts": [
            str(json_path.relative_to(SOURCE_FILE.parent.parent)),
            str(jsonl_path.relative_to(SOURCE_FILE.parent.parent)),
            str(manifest_path.relative_to(SOURCE_FILE.parent.parent)),
        ],
    }
    with open(manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, ensure_ascii=False, indent=2)
    print(f"   ✅ {manifest_path.stat().st_size:,} bytes")

    # ── Step 8: Validation ────────────────────────────────────────────────
    print()
    print("🔬 Step 8: Validation...")
    errors = 0

    # Validate JSON roundtrip
    with open(json_path, "r", encoding="utf-8") as f:
        loaded = json.load(f)
    if loaded["total_records"] == len(unique_records):
        print("   ✅ JSON roundtrip: PASS")
    else:
        print("   ❌ JSON roundtrip: FAIL")
        errors += 1

    # Validate JSONL line count
    with open(jsonl_path, "r", encoding="utf-8") as f:
        line_count = sum(1 for _ in f)
    if line_count == len(unique_records):
        print("   ✅ JSONL line count: PASS")
    else:
        print(f"   ❌ JSONL line count: FAIL (expected {len(unique_records)}, got {line_count})")
        errors += 1

    # Validate all records have required fields
    required_fields = {"id", "program_scope", "category", "content", "source_file", "source_sheet", "content_sha256"}
    missing = []
    for rec in unique_records:
        for field in required_fields:
            if field not in rec or rec[field] is None:
                missing.append((rec["id"], field))
    if not missing:
        print("   ✅ Required fields: PASS")
    else:
        print(f"   ⚠️  Required fields: {len(missing)} missing values")
        for rec_id, field in missing[:5]:
            print(f"      {rec_id}: missing '{field}'")

    # Validate unique IDs
    ids = [r["id"] for r in unique_records]
    if len(ids) == len(set(ids)):
        print("   ✅ Unique IDs: PASS")
    else:
        print("   ❌ Unique IDs: FAIL (duplicates found)")
        errors += 1

    # ── Summary ───────────────────────────────────────────────────────────
    print()
    print("╔══════════════════════════════════════════════════════════════╗")
    print("║                    CONVERSION COMPLETE                      ║")
    print("╠══════════════════════════════════════════════════════════════╣")
    print(f"║  Source:     {SOURCE_FILE.name:46s} ║")
    print(f"║  Records:   {len(unique_records):<46d} ║")
    print(f"║  Unique:    {unique_count:<46d} ║")
    print(f"║  Duplicates: {dup_count:<45d} ║")
    print(f"║  Sheets:    {len(wb.sheetnames):<46d} ║")
    print(f"║  Errors:    {errors:<46d} ║")
    print("╠══════════════════════════════════════════════════════════════╣")
    print(f"║  📁 {str(json_path):55s} ║")
    print(f"║  📁 {str(jsonl_path):55s} ║")
    print(f"║  📁 {str(manifest_path):55s} ║")
    print("╚══════════════════════════════════════════════════════════════╝")

    return 0 if errors == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
